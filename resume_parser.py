"""
resume_parser.py
Module for extracting raw text from PDF and DOCX resume documents.
Gracefully handles file formats, corrupted files, and empty contents.
"""

import io
from typing import Optional, Tuple, Union


def extract_text_from_pdf(file_source: Union[str, io.BytesIO, bytes]) -> Tuple[str, Optional[str]]:
    """
    Extract text from a PDF file using pdfplumber with PyPDF2/pypdf as fallback.

    Args:
        file_source: File path, BytesIO stream, or raw bytes.

    Returns:
        Tuple of (extracted_text, error_message)
    """
    text_content = []
    error = None

    # Normalize file_source into a stream if it is bytes or a path string
    if isinstance(file_source, bytes):
        stream = io.BytesIO(file_source)
    elif isinstance(file_source, str):
        try:
            with open(file_source, "rb") as f:
                stream = io.BytesIO(f.read())
        except Exception as e:
            return "", f"Could not open file: {str(e)}"
    else:
        stream = file_source

    # Primary approach: pdfplumber (best for layout and table preservation)
    try:
        import pdfplumber

        # Ensure pointer is at beginning
        if hasattr(stream, "seek"):
            stream.seek(0)

        with pdfplumber.open(stream) as pdf:
            if not pdf.pages:
                return "", "The uploaded PDF has no pages."

            for page_idx, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text()
                if page_text:
                    text_content.append(page_text.strip())

        extracted = "\n\n".join(text_content).strip()
        if extracted:
            return extracted, None
    except Exception as plumber_err:
        error = f"pdfplumber extraction failed ({str(plumber_err)})"

    # Fallback approach: pypdf or PyPDF2
    try:
        try:
            import pypdf as pypdf_module
        except ImportError:
            import PyPDF2 as pypdf_module

        if hasattr(stream, "seek"):
            stream.seek(0)

        reader = pypdf_module.PdfReader(stream)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                return "", "PDF is password-protected and cannot be read."

        fallback_text = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                fallback_text.append(t.strip())

        extracted = "\n\n".join(fallback_text).strip()
        if extracted:
            return extracted, None
        elif not text_content:
            return "", "The PDF appears to be empty or contains scanned images without text."
    except Exception as pypdf_err:
        return "", f"Failed to extract text from PDF: {error or str(pypdf_err)}"

    return "", "Could not extract text from the PDF file."


def extract_text_from_docx(file_source: Union[str, io.BytesIO, bytes]) -> Tuple[str, Optional[str]]:
    """
    Extract text from a DOCX file using python-docx.
    Extracts text from paragraphs as well as embedded tables.

    Args:
        file_source: File path, BytesIO stream, or raw bytes.

    Returns:
        Tuple of (extracted_text, error_message)
    """
    try:
        import docx
    except ImportError:
        return "", "python-docx library is not installed."

    if isinstance(file_source, bytes):
        stream = io.BytesIO(file_source)
    elif isinstance(file_source, str):
        try:
            with open(file_source, "rb") as f:
                stream = io.BytesIO(f.read())
        except Exception as e:
            return "", f"Could not open file: {str(e)}"
    else:
        stream = file_source

    try:
        if hasattr(stream, "seek"):
            stream.seek(0)

        doc = docx.Document(stream)
        text_parts = []

        # 1. Extract from paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                text_parts.append(text)

        # 2. Extract from tables (many resumes place skills in multi-column tables)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Deduplicate repeated cell texts from merged cells
                    unique_cells = []
                    for c in row_cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    text_parts.append(" | ".join(unique_cells))

        extracted = "\n".join(text_parts).strip()
        if not extracted:
            return "", "The DOCX file is empty or contains no readable text."

        return extracted, None
    except Exception as e:
        return "", f"Error reading DOCX document: {str(e)}"


def parse_resume(uploaded_file) -> Tuple[str, Optional[str]]:
    """
    High-level entry point to parse a resume from an uploaded file or file path.

    Args:
        uploaded_file: Streamlit UploadedFile, file-like object, or file path.

    Returns:
        Tuple of (extracted_text, error_message)
    """
    if uploaded_file is None:
        return "", "No file was provided."

    # Determine filename
    filename = getattr(uploaded_file, "name", None)
    if not filename and isinstance(uploaded_file, str):
        filename = uploaded_file

    if not filename:
        return "", "Unknown file format. Please upload a PDF or DOCX file."

    filename_lower = filename.lower()

    # Check file size if available
    size = getattr(uploaded_file, "size", None)
    if size == 0:
        return "", "The uploaded file is empty (0 bytes)."

    if filename_lower.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    elif filename_lower.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    elif filename_lower.endswith(".doc"):
        return "", ".doc legacy format is not directly supported. Please convert and upload as .docx or .pdf."
    else:
        return "", f"Unsupported file type: '{filename}'. Please upload a PDF (.pdf) or Word document (.docx)."
