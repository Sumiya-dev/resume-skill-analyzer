"""
generate_test_resumes.py
Generates realistic sample PDF and DOCX resumes in sample_resumes/
for testing and demonstration purposes.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_docx_resume(filepath: str):
    doc = docx.Document()
    
    # Header
    title = doc.add_heading("Alex Mercer", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph("Full-Stack Software Engineer | San Francisco, CA | alex.mercer@email.com | (555) 019-2834")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Summary
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(
        "Accomplished Software Engineer with 5+ years of experience in architecting, developing, and deploying "
        "enterprise web applications and distributed systems. Adept in modern JavaScript frameworks, Python backend "
        "development, RESTful API design, and cloud container orchestration."
    )
    
    # Skills Table
    doc.add_heading("Technical Skills", level=1)
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Technical Domain"
    hdr_cells[1].text = "Skills & Frameworks"
    
    skills = [
        ("Programming Languages", "Python, Java, TypeScript, JavaScript, C++, C#, SQL, HTML, CSS, Bash"),
        ("Web Development", "React, React.js, Next.js, Node.js, Express.js, FastAPI, Django, Flask, Tailwind CSS, GraphQL, REST API, Redux"),
        ("Databases", "PostgreSQL, MongoDB, Redis, MySQL, SQLite, Supabase"),
        ("AI / Machine Learning", "Machine Learning, Deep Learning, Natural Language Processing, NLP, PyTorch, TensorFlow, Scikit-learn, Hugging Face"),
        ("Data Science", "Pandas, NumPy, Matplotlib, Seaborn, Power BI, Statistics"),
        ("Tools & Platforms", "Git, GitHub, VS Code, Linux, Ubuntu, Docker, Postman, Jira, Figma"),
        ("Cloud & DevOps", "AWS, Microsoft Azure, Kubernetes, CI/CD, GitHub Actions, Terraform, Nginx")
    ]
    
    for domain, skill_str in skills:
        row_cells = table.add_row().cells
        row_cells[0].text = domain
        row_cells[1].text = skill_str
        
    # Experience
    doc.add_heading("Work Experience", level=1)
    
    doc.add_heading("Senior Full-Stack Engineer — CloudNova Technologies (2022 - Present)", level=2)
    doc.add_paragraph(
        "• Engineered microservices with Python (FastAPI) and Node.js serving 2.5M requests daily.\n"
        "• Built interactive user dashboards using React.js, Next.js, and Tailwind CSS.\n"
        "• Optimized PostgreSQL database queries and implemented Redis caching to reduce latency by 45%.\n"
        "• Designed and maintained CI/CD deployment pipelines using Docker, Kubernetes, and AWS."
    )
    
    doc.add_heading("Software Engineer — Nexus Systems (2020 - 2022)", level=2)
    doc.add_paragraph(
        "• Developed scalable backend endpoints in Django and REST API services.\n"
        "• Collaborated on cross-platform features using TypeScript and Redux state management.\n"
        "• Automated unit testing and integrated GitHub Actions workflows."
    )
    
    # Education
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. in Computer Science — University of California, Berkeley (2016 - 2020)")
    
    doc.save(filepath)
    print(f"Created DOCX resume: {filepath}")


def create_empty_docx(filepath: str):
    doc = docx.Document()
    doc.add_heading("Meeting Notes and Shopping List", level=1)
    doc.add_paragraph("This document contains general administrative notes without any technical skill keywords.")
    doc.add_paragraph("Agenda: Review quarterly budget, order supplies, schedule team offsite.")
    doc.save(filepath)
    print(f"Created empty DOCX: {filepath}")


def create_pdf_resume(filepath: str):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    story = []

    # Title style
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1e293b'),
        alignment=1, # Center
        spaceAfter=4
    )
    
    sub_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#64748b'),
        alignment=1,
        spaceAfter=14
    )
    
    h1_style = ParagraphStyle(
        'SectionH1',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#0284c7'),
        spaceBefore=10,
        spaceAfter=6
    )
    
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    )

    story.append(Paragraph("Sarah Jenkins", title_style))
    story.append(Paragraph("AI & Data Science Specialist | New York, NY | sarah.j@example.com | (555) 987-6543", sub_style))

    story.append(Paragraph("Professional Summary", h1_style))
    story.append(Paragraph(
        "Data Scientist and Machine Learning Engineer with 4+ years of expertise in NLP, Computer Vision, "
        "and predictive modeling. Strong hands-on background in PyTorch, TensorFlow, Hugging Face Transformers, "
        "and big data analytics with Spark and SQL.",
        body_style
    ))

    story.append(Paragraph("Technical Skills", h1_style))
    skills_data = [
        [Paragraph("<b>Programming Languages</b>", body_style), Paragraph("Python, R, SQL, C++, Bash", body_style)],
        [Paragraph("<b>AI / Machine Learning</b>", body_style), Paragraph("Machine Learning, Deep Learning, Natural Language Processing, NLP, PyTorch, TensorFlow, Keras, Scikit-learn, Hugging Face, Transformers, OpenCV, BERT, LLMs", body_style)],
        [Paragraph("<b>Data Science</b>", body_style), Paragraph("Pandas, NumPy, SciPy, Matplotlib, Seaborn, Power BI, Tableau, Statistics, Exploratory Data Analysis, EDA, Apache Spark, PySpark", body_style)],
        [Paragraph("<b>Databases</b>", body_style), Paragraph("PostgreSQL, SQLite, MongoDB, BigQuery", body_style)],
        [Paragraph("<b>Cloud / DevOps</b>", body_style), Paragraph("AWS, Google Cloud, Docker, Kubernetes, CI/CD", body_style)],
        [Paragraph("<b>Tools / Platforms</b>", body_style), Paragraph("Git, GitHub, Linux, Jupyter Notebook, VS Code", body_style)]
    ]
    
    table = Table(skills_data, colWidths=[150, 380])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(table)
    story.append(Spacer(1, 10))

    story.append(Paragraph("Experience", h1_style))
    story.append(Paragraph("<b>Senior NLP Engineer — Synthetix AI (2022 - Present)</b>", body_style))
    story.append(Paragraph(
        "• Fine-tuned large language models (LLMs) with Hugging Face and PyTorch for domain-specific text summarization.<br/>"
        "• Designed automated resume screening pipelines utilizing spaCy and rule-based token matchers.<br/>"
        "• Built ETL pipelines with Pandas and PySpark processing over 500GB daily on AWS.",
        body_style
    ))
    
    story.append(Paragraph("Education", h1_style))
    story.append(Paragraph("M.S. in Data Science — Columbia University | B.S. in Statistics — NYU", body_style))

    doc.build(story)
    print(f"Created PDF resume: {filepath}")


def create_software_engineer_pdf(filepath: str):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    doc = SimpleDocTemplate(filepath, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=20, leading=24, textColor=colors.HexColor('#0f172a'), alignment=1, spaceAfter=4)
    sub_style = ParagraphStyle('DocSubtitle', parent=styles['Normal'], fontSize=10, leading=13, textColor=colors.HexColor('#64748b'), alignment=1, spaceAfter=14)
    h1_style = ParagraphStyle('SectionH1', parent=styles['Heading2'], fontSize=13, leading=16, textColor=colors.HexColor('#2563eb'), spaceBefore=10, spaceAfter=6)
    body_style = ParagraphStyle('BodyDark', parent=styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor('#334155'), spaceAfter=6)

    story.append(Paragraph("Alex Mercer", title_style))
    story.append(Paragraph("Full-Stack Software Engineer | San Francisco, CA | alex.mercer@email.com | (555) 019-2834", sub_style))

    story.append(Paragraph("Professional Summary", h1_style))
    story.append(Paragraph(
        "Accomplished Software Engineer with 5+ years of experience in architecting, developing, and deploying "
        "enterprise web applications and distributed systems. Adept in modern JavaScript frameworks, Python backend "
        "development, RESTful API design, and cloud container orchestration.",
        body_style
    ))

    story.append(Paragraph("Technical Skills", h1_style))
    skills_data = [
        [Paragraph("<b>Programming Languages</b>", body_style), Paragraph("Python, Java, TypeScript, JavaScript, C++, C#, SQL, HTML, CSS, Bash", body_style)],
        [Paragraph("<b>Web Development</b>", body_style), Paragraph("React, React.js, Next.js, Node.js, Express.js, FastAPI, Django, Flask, Tailwind CSS, GraphQL, REST API, Redux", body_style)],
        [Paragraph("<b>Databases</b>", body_style), Paragraph("PostgreSQL, MongoDB, Redis, MySQL, SQLite, Supabase", body_style)],
        [Paragraph("<b>AI / Machine Learning</b>", body_style), Paragraph("Machine Learning, Deep Learning, Natural Language Processing, NLP, PyTorch, TensorFlow, Scikit-learn, Hugging Face", body_style)],
        [Paragraph("<b>Data Science</b>", body_style), Paragraph("Pandas, NumPy, Matplotlib, Seaborn, Power BI, Statistics", body_style)],
        [Paragraph("<b>Tools & Platforms</b>", body_style), Paragraph("Git, GitHub, VS Code, Linux, Ubuntu, Docker, Postman, Jira, Figma", body_style)],
        [Paragraph("<b>Cloud & DevOps</b>", body_style), Paragraph("AWS, Microsoft Azure, Kubernetes, CI/CD, GitHub Actions, Terraform, Nginx", body_style)]
    ]

    table = Table(skills_data, colWidths=[150, 380])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor('#334155')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(table)
    story.append(Spacer(1, 10))

    story.append(Paragraph("Work Experience", h1_style))
    story.append(Paragraph("<b>Senior Full-Stack Engineer — CloudNova Technologies (2022 - Present)</b>", body_style))
    story.append(Paragraph(
        "• Engineered microservices with Python (FastAPI) and Node.js serving 2.5M requests daily.<br/>"
        "• Built interactive user dashboards using React.js, Next.js, and Tailwind CSS.<br/>"
        "• Optimized PostgreSQL database queries and implemented Redis caching to reduce latency by 45%.<br/>"
        "• Designed and maintained CI/CD deployment pipelines using Docker, Kubernetes, and AWS.",
        body_style
    ))
    story.append(Paragraph("<b>Software Engineer — Nexus Systems (2020 - 2022)</b>", body_style))
    story.append(Paragraph(
        "• Developed scalable backend endpoints in Django and REST API services.<br/>"
        "• Collaborated on cross-platform features using TypeScript and Redux state management.<br/>"
        "• Automated unit testing and integrated GitHub Actions workflows.",
        body_style
    ))

    story.append(Paragraph("Education", h1_style))
    story.append(Paragraph("B.S. in Computer Science — University of California, Berkeley (2016 - 2020)", body_style))

    doc.build(story)
    print(f"Created PDF resume: {filepath}")


def create_devops_engineer_pdf(filepath: str):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    doc = SimpleDocTemplate(filepath, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=20, leading=24, textColor=colors.HexColor('#0f172a'), alignment=1, spaceAfter=4)
    sub_style = ParagraphStyle('DocSubtitle', parent=styles['Normal'], fontSize=10, leading=13, textColor=colors.HexColor('#64748b'), alignment=1, spaceAfter=14)
    h1_style = ParagraphStyle('SectionH1', parent=styles['Heading2'], fontSize=13, leading=16, textColor=colors.HexColor('#0891b2'), spaceBefore=10, spaceAfter=6)
    body_style = ParagraphStyle('BodyDark', parent=styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor('#334155'), spaceAfter=6)

    story.append(Paragraph("David Clark", title_style))
    story.append(Paragraph("Senior Infrastructure & Cloud Engineer | Austin, TX | david.c@example.com | (555) 345-6789", sub_style))

    story.append(Paragraph("Core Expertise", h1_style))
    story.append(Paragraph(
        "Senior Infrastructure & Cloud Engineer with 7+ years of experience specializing in infrastructure automation, "
        "cloud architecture, continuous integration, and container orchestration across multi-cloud environments.",
        body_style
    ))

    story.append(Paragraph("Technical Skills", h1_style))
    skills_data = [
        [Paragraph("<b>Cloud & DevOps</b>", body_style), Paragraph("AWS, Microsoft Azure, Google Cloud Platform, GCP, Docker, Kubernetes, Terraform, Ansible, Jenkins, CI/CD, GitHub Actions, Helm, Prometheus, Grafana, Nginx", body_style)],
        [Paragraph("<b>Programming Languages</b>", body_style), Paragraph("Python, Go, Golang, Bash, Shell Scripting, PowerShell, SQL", body_style)],
        [Paragraph("<b>Databases</b>", body_style), Paragraph("MySQL, PostgreSQL, Redis, Elasticsearch", body_style)],
        [Paragraph("<b>Tools & Platforms</b>", body_style), Paragraph("Git, GitLab, Linux, Ubuntu, Jira, Confluence, Slack", body_style)]
    ]

    table = Table(skills_data, colWidths=[150, 380])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor('#334155')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(table)
    story.append(Spacer(1, 10))

    story.append(Paragraph("Work Experience", h1_style))
    story.append(Paragraph("<b>Lead DevOps Engineer — ScaleOps (2021 - Present)</b>", body_style))
    story.append(Paragraph(
        "• Provisioned multi-region Kubernetes clusters on AWS managed via Terraform and Helm.<br/>"
        "• Reduced deployment pipeline cycle from 40 minutes to under 5 minutes using GitHub Actions.<br/>"
        "• Implemented end-to-end monitoring and observability stacks with Prometheus and Grafana.",
        body_style
    ))

    story.append(Paragraph("Education", h1_style))
    story.append(Paragraph("B.S. in Information Systems — University of Texas at Austin", body_style))

    doc.build(story)
    print(f"Created PDF resume: {filepath}")


def create_empty_pdf(filepath: str):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet

    doc = SimpleDocTemplate(filepath, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Company Newsletter — October Issue", styles['Heading1']),
        Paragraph("This is an internal newsletter containing general event details and announcements with no resume keywords.", styles['Normal'])
    ]
    doc.build(story)
    print(f"Created empty PDF: {filepath}")


def main():
    out_dir = "sample_resumes"
    os.makedirs(out_dir, exist_ok=True)
    
    create_docx_resume(os.path.join(out_dir, "sample_software_engineer_resume.docx"))
    create_software_engineer_pdf(os.path.join(out_dir, "sample_software_engineer_resume.pdf"))
    create_empty_docx(os.path.join(out_dir, "empty_sample_resume.docx"))
    create_pdf_resume(os.path.join(out_dir, "sample_ai_engineer_resume.pdf"))
    create_devops_engineer_pdf(os.path.join(out_dir, "sample_devops_engineer_resume.pdf"))
    create_empty_pdf(os.path.join(out_dir, "empty_sample_resume.pdf"))
    print("All sample resumes successfully generated in sample_resumes/ directory!")


if __name__ == "__main__":
    main()
